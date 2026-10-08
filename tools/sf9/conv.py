import openpyxl, html, json, re, sys
from openpyxl.utils import get_column_letter as L

MDW = 7.0          # px per width unit (Arial 10 is the workbook's default font)
PT = 96/72

WEIGHT = {'thick':4,'medium':3,'thin':2,'double':3,'dashed':1,'dotted':1,'hair':1,'mediumDashed':2,'dashDot':1}
def css_border(side):
    s = side.style if side is not None else None
    if not s: return None
    w = {'thick':'3px','medium':'2px','double':'3px'}.get(s,'1px')
    st = {'double':'double','dotted':'dotted','hair':'dotted','dashed':'dashed','mediumDashed':'dashed','dashDot':'dashed'}.get(s,'solid')
    return f'{w} {st} #000'

def color_of(c, default=None):
    if c is None: return default
    if c.type == 'rgb' and isinstance(c.rgb, str) and len(c.rgb) == 8:
        return '#' + c.rgb[2:]
    if c.type == 'theme':
        t = c.tint or 0
        base = {0:255, 1:0}.get(c.theme)
        if base is None: return default
        v = base*(1+t) if t<0 else base+(255-base)*t
        v = int(round(v)); return f'#{v:02x}{v:02x}{v:02x}'
    if c.type == 'indexed' and c.indexed == 64: return default
    return default

FONT_FALLBACK = '"Bookman Old Style","Bookman","URW Bookman","Palatino Linotype",Georgia,serif'

class Sheet:
    def __init__(self, path, c1, c2, r1, r2):
        self.wb = openpyxl.load_workbook(path); self.ws = self.wb.active
        ws = self.ws
        self.w = {}
        for k, v in ws.column_dimensions.items():
            for i in range(v.min, v.max+1): self.w[i] = v.width
        self.merge = {}      # anchor (r,c) -> (r2,c2)
        self.covered = set()
        for m in ws.merged_cells.ranges:
            self.merge[(m.min_row, m.min_col)] = (m.max_row, m.max_col)
            for r in range(m.min_row, m.max_row+1):
                for c in range(m.min_col, m.max_col+1):
                    if (r,c) != (m.min_row, m.min_col): self.covered.add((r,c))
        self.c1, self.c2, self.r1, self.r2 = c1, c2, r1, r2
        self.styles = {}     # key -> class
    def colpx(self, c): return round((self.w.get(c) or ws_default) * MDW) if False else round((self.w.get(c) or 12.5703125) * MDW)
    def rowpx(self, r):
        h = self.ws.row_dimensions[r].height
        return round((h if h else 15.0) * PT, 1)
    def region(self, r, c):
        a = self.merge.get((r,c))
        return (a[0], a[1]) if a else (r, c)
    def edge(self, r, c, side):
        """Strongest border on this cell's (or its merged block's) outer edge, taking the neighbour's facing edge too."""
        ws = self.ws
        r2, c2 = self.region(r, c)
        best, bw = None, 0
        def consider(rr, cc, sd):
            nonlocal best, bw
            if rr < 1 or cc < 1: return
            b = getattr(ws.cell(rr, cc).border, sd)
            if b is not None and b.style and WEIGHT.get(b.style, 1) > bw:
                best, bw = b, WEIGHT.get(b.style, 1)
        if side == 'top':
            for cc in range(c, c2+1): consider(r, cc, 'top'); consider(r-1, cc, 'bottom')
        elif side == 'bottom':
            for cc in range(c, c2+1): consider(r2, cc, 'bottom'); consider(r2+1, cc, 'top')
        elif side == 'left':
            for rr in range(r, r2+1): consider(rr, c, 'left'); consider(rr, c-1, 'right')
        else:
            for rr in range(r, r2+1): consider(rr, c2, 'right'); consider(rr, c2+1, 'left')
        return css_border(best) if best is not None else None
    def style_class(self, r, c, extra=None):
        ws = self.ws; cell = ws.cell(r, c); f = cell.font; al = cell.alignment
        props = []
        fam = f.name or 'Arial'
        props.append('font-family:' + (FONT_FALLBACK if 'Bookman' in fam else f'"{fam}",Arial,sans-serif'))
        props.append(f'font-size:{round((f.sz or 10)*PT,2)}px')
        if f.b: props.append('font-weight:700')
        if f.i: props.append('font-style:italic')
        if f.u: props.append('text-decoration:underline')
        col = color_of(f.color, '#000000')
        if col != '#000000': props.append('color:' + col)
        h = al.horizontal
        if h in ('center','centerContinuous'): props.append('text-align:center')
        elif h == 'right': props.append('text-align:right')
        else: props.append('text-align:left')
        v = al.vertical
        props.append('vertical-align:' + {'center':'middle','top':'top'}.get(v, 'bottom'))
        if al.wrap_text: props.append('white-space:pre-wrap;word-break:break-word')
        else: props.append('white-space:pre;overflow:visible')
        if al.indent: props.append(f'padding-left:{int(al.indent*9)+2}px')
        fill = cell.fill
        if fill and fill.fill_type == 'solid':
            bg = color_of(fill.fgColor, None)
            if bg and bg.lower() != '#ffffff': props.append('background:' + bg)
        for side in ('top','right','bottom','left'):
            b = self.edge(r, c, side)
            props.append(f'border-{side}:{b if b else "none"}')
        key = ';'.join(props)
        if key not in self.styles: self.styles[key] = None
        return key
    def row_html(self, r, roles, c1=None, c2=None, collapse_rows=False, extra_h=0, role_attr=True):
        c1 = c1 or self.c1; c2 = c2 or self.c2
        ws = self.ws
        cells = []
        c = c1
        while c <= c2:
            if (r, c) in self.covered: c += 1; continue
            m = self.merge.get((r, c))
            rs, cs = 1, 1
            if m:
                rs = 1 if collapse_rows else (m[0]-r+1)
                cs = min(m[1], c2) - c + 1
            cell = ws.cell(r, c)
            val = cell.value
            coord = f'{L(c)}{r}'
            role = roles.get(coord)
            text = ''
            if val is not None and not (isinstance(val, str) and val.startswith('=')):
                text = str(val)
            if role and role.startswith('!'): text = ''; role = role[1:]
            if role == 'keep': role = None
            key = self.style_class(r, c)
            attrs = f' class="@@{key}@@"'
            if cs > 1: attrs += f' colspan="{cs}"'
            if rs > 1: attrs += f' rowspan="{rs}"'
            if role: attrs += f' data-f="{role}"'
            t = html.escape(text).replace('\n', '<br>')
            cells.append(f'<td{attrs}>{t}</td>')
            c += cs if cs > 1 else 1
        h = self.rowpx(r) + extra_h
        return f'<tr style="height:{h}px">' + ''.join(cells) + '</tr>'
    def colgroup(self, c1, c2):
        return '<colgroup>' + ''.join(f'<col style="width:{self.colpx(c)}px">' for c in range(c1, c2+1)) + '</colgroup>', sum(self.colpx(c) for c in range(c1, c2+1))

# Which cells the app fills, by role. Everything else prints exactly as the template has it.
ROLES = {
  'B7':'region','B8':'division','B9':'district','B10':'municipality','B12':'school','B15':'year',
  'C17':'name','I17':'age','K17':'sex','C18':'lrn','I18':'grade','K18':'section','C19':'track',
  'B27':'principal','H27':'adviser',
}
CD = ['R','S','T','U','V','W','X','Y','Z','AA','AB']
for i, col in enumerate(CD): ROLES[f'{col}5'] = f'cd{i}'
ROLES['AC5'] = 'cdtotal'
for col in CD + ['AC']:
    for rr in (7, 9): ROLES[f'{col}{rr}'] = 'keep'

def build(path, spec):
    S = Sheet(path, 1, 12, 1, 52)
    ws = S.ws
    roles = dict(ROLES)
    # the formula cells in the right panel that repeat the adviser / school head
    for row in ws.iter_rows(min_row=1, max_row=54, min_col=16, max_col=30):
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in ('=H27','=B27','=Q41'):
                roles[c.coordinate] = 'principal' if c.value.strip() in ('=B27','=Q41') else 'adviser'
    # label cells whose template text is a placeholder
    roles['B7'] = 'region'; roles['B8'] = 'division'; roles['B9'] = 'district'; roles['B10'] = 'municipality'
    # coordinates of the value cells that hold a template literal need blanking when replaced by role: handled in JS
    out = {}
    lcg, lw = S.colgroup(1, 12)
    rcg, rw = S.colgroup(16, 30)
    out['leftCols'] = lcg; out['leftW'] = lw; out['rightCols'] = rcg; out['rightW'] = rw
    gap = sum(S.colpx(c) for c in range(13, 16)); out['gapW'] = gap
    h1, h2 = spec['head']
    out['leftHead'] = ''.join(S.row_html(r, roles, 1, 12) for r in range(h1, h2+1))
    t1, t2 = spec['tail']
    out['leftTail'] = ''.join(S.row_html(r, roles, 1, 12) for r in range(t1, t2+1))
    tpl = {}
    for name, rr in spec['rows'].items():
        if name == 'ga':
            extra = sum(S.rowpx(rr+i) for i in range(1, spec.get('gaRows', 1)))
            tpl[name] = S.row_html(rr, {}, 1, 12, collapse_rows=True, extra_h=extra)
        else:
            tpl[name] = S.row_html(rr, {}, 1, 12)
    out['rows'] = tpl
    rr1, rr2 = spec['right']
    out['right'] = ''.join(S.row_html(r, roles, 16, 30) for r in range(rr1, rr2+1))
    out['gray'] = spec.get('gray')
    return S, out

def finish(frames, sheets):
    # assign short class names across both frames
    keys = {}
    for S in sheets:
        for k in S.styles: keys.setdefault(k, None)
    names = {k: f's{i}' for i, k in enumerate(keys)}
    css = ''.join(f'.sf9 .{names[k]}{{{k}}}' for k in keys)
    for f in frames.values():
        def sub(s): return re.sub(r'@@(.*?)@@', lambda m: names[m.group(1)], s)
        for k in ('leftHead','leftTail','right'): f[k] = sub(f[k])
        f['rows'] = {a: sub(b) for a, b in f['rows'].items()}
    return css

if __name__ == '__main__':
    U = './templates/'  # put the SF9 .xlsx templates here
    specs = {
      'k10': (U+'SF_9_-_Grades_4-10.xlsx', dict(head=(2,31), tail=(44,52), right=(2,52), rows={'subject':33,'component':40,'ga':42}, gaRows=2)),
      'shs': (U+'SF_9_-_Grade_11_-_Academic.xlsx', dict(head=(2,31), tail=(45,53), right=(2,53), rows={'category':32,'subject':36,'component':34,'ga':44}, gaRows=1, gray='#a5a5a5')),
    }
    frames = {}; sheets = []
    for name, (path, spec) in specs.items():
        S, out = build(path, spec); frames[name] = out; sheets.append(S)
    css = finish(frames, sheets)
    # Drawings (same anchors in every template): the DepEd seal and the "Insert School Logo" circle, in px from the
    # left panel's top-left corner (the panel starts at sheet row 2).
    S0 = sheets[0]
    E = 9525.0
    top_of = lambda row0: sum(S0.rowpx(r) for r in range(2, row0 + 1))   # row0 is the 0-based row of the anchor
    left_of = lambda col0: sum(S0.colpx(c) for c in range(1, col0 + 1))
    place = {
      'seal': {'x': round(247650/E + left_of(0)), 'y': round(top_of(3)), 'size': round(1314450/E)},
      'logo': {'x': round(438150/E + left_of(9)), 'y': round(top_of(2) + 190500/E), 'w': round(1276350/E), 'h': round(1314450/E)},
    }
    seal = open('./seal.b64').read()
    data = {'css': css, 'frames': frames, 'place': place, 'seal': 'data:image/png;base64,' + seal}
    json.dump(data, open('./frames.json','w'))
    open('./sf9-frames.js','w').write('// GENERATED from the DepEd SF9 Excel templates by sf9/conv.py -- do not edit by hand.\nvar SF9_DATA = ' + json.dumps(data, separators=(',',':')) + ';\n')
    print('place', place)
    print('css bytes', len(css), {k: {kk: (len(vv) if isinstance(vv,str) else vv) for kk, vv in v.items() if kk!='rows'} for k, v in frames.items()})
